"""导出待审核清单的测试（RED → GREEN）。

不触碰文件系统：索引用内存 SQLite，工作簿用 BytesIO 往返。
验证范围过滤、列顺序、下拉校验、位置可读化、序号与说明页。
"""

import io
import json
import sqlite3
import unittest

from openpyxl import load_workbook

from export_review_sheet import (
    JUDGE_OPTIONS,
    POPULATION_OPTIONS,
    REVIEW_SCOPE_TITLES,
    build_review_workbook,
    human_locator,
    load_chunks,
)


def _meta(title, chunk_id, locator, text, parent="父段落", pdf_page=None, sha="a" * 64):
    return {
        "title": title,
        "chunk_id": chunk_id,
        "locator": locator,
        "section": "待审核章节",
        "pdf_page": pdf_page,
        "printed_page": None,
        "text": text,
        "parent_text": parent,
        "population": "unknown",
        "care_stage": "unknown",
        "layout_status": "pending_review",
        "document_sha256": sha,
        "raw_path": "raw/demo/" + title,
    }


class ReviewSheetExportTest(unittest.TestCase):
    def setUp(self):
        self.con = sqlite3.connect(":memory:")
        self.addCleanup(self.con.close)
        self.con.execute("CREATE TABLE chunks (id TEXT, metadata TEXT, vector BLOB)")
        payload = [
            _meta("NICU出院高危儿0～3岁生长发育随访管理技术的专家共识 (1)", "c1", "paragraph:3", "示例一"),
            _meta("NICU出院高危儿0～3岁生长发育随访管理技术的专家共识 (1)", "c2", "pdf:2:lines:1-4", "示例二", pdf_page=2),
            _meta("语料库（终版）", "c3", "paragraph:1", "汇编片段"),
            _meta("特殊状态儿童预防接种(广东)专家共识", "c4", "pdf:5:lines:2-9", "示例四", pdf_page=5),
            _meta("早产儿保健服务指南", "c5", "paragraph:9", "与 PDF 通知重复的附件"),
        ]
        self.con.executemany(
            "INSERT INTO chunks (id, metadata) VALUES (?, ?)",
            [(f"row{i}", json.dumps(m, ensure_ascii=False)) for i, m in enumerate(payload)],
        )
        self.con.commit()

    def _rows(self):
        return load_chunks(titles=REVIEW_SCOPE_TITLES, connection=self.con)

    def test_只导出范围内文档并排除汇编与重复文档(self):
        titles = {r["title"] for r in self._rows()}
        self.assertIn("NICU出院高危儿0～3岁生长发育随访管理技术的专家共识 (1)", titles)
        self.assertIn("特殊状态儿童预防接种(广东)专家共识", titles)
        self.assertNotIn("语料库（终版）", titles)
        # 与 PDF 通知逐字重复的 DOCX 附件默认不导出
        self.assertNotIn("早产儿保健服务指南", titles)
        self.assertEqual(len(self._rows()), 3)

    def test_显式要求时才导出重复文档(self):
        rows = load_chunks(
            titles=REVIEW_SCOPE_TITLES, connection=self.con, include_duplicates=True
        )
        self.assertIn("早产儿保健服务指南", {r["title"] for r in rows})

    def test_位置转成可读文本(self):
        self.assertEqual(
            human_locator({"locator": "pdf:2:lines:1-4", "pdf_page": 2}), "PDF 第 2 页 · 行 1-4"
        )
        self.assertEqual(human_locator({"locator": "paragraph:3"}), "DOCX 第 3 段")
        self.assertEqual(human_locator({"locator": ""}), "位置未知")

    def test_工作簿结构符合约定(self):
        ws = build_review_workbook(self._rows(), {"run_id": "run-x"})["待审核清单"]
        self.assertEqual(
            [c.value for c in ws[1]],
            [
                "序号", "片段ID（勿改）", "文档", "类型", "原文位置", "页", "章节",
                "片段文本", "父段落", "字数", "质量提示", "判定", "适用人群",
                "护理场景", "修改建议", "审核人", "审核日期", "备注",
            ],
        )
        self.assertEqual([ws.cell(row=r, column=1).value for r in range(2, 5)], [1, 2, 3])
        self.assertIsNone(ws.cell(row=5, column=1).value)
        self.assertEqual(
            {ws.cell(row=r, column=2).value for r in range(2, 5)}, {"c1", "c2", "c4"}
        )

    def test_判定与元数据列有下拉校验(self):
        ws = build_review_workbook(self._rows(), {})["待审核清单"]
        refs = {str(dv.sqref): dv for dv in ws.data_validations.dataValidation}
        self.assertTrue(any(ref.startswith("L2") for ref in refs), "判定列缺少下拉")
        judge = next(dv for ref, dv in refs.items() if ref.startswith("L2"))
        self.assertIn("通过", judge.formula1)
        self.assertTrue(any(ref.startswith("M2") for ref in refs), "人群列缺少下拉")
        self.assertTrue(any(ref.startswith("N2") for ref in refs), "场景列缺少下拉")
        self.assertIn("通过", JUDGE_OPTIONS)
        self.assertIn("早产儿", POPULATION_OPTIONS)

    def test_同文档内按页与行号排序(self):
        con = sqlite3.connect(":memory:")
        self.addCleanup(con.close)
        con.execute("CREATE TABLE chunks (id TEXT, metadata TEXT, vector BLOB)")
        title = "NICU出院高危儿0～3岁生长发育随访管理技术的专家共识 (1)"
        payload = [
            _meta(title, "aaa-late", "pdf:1:lines:8-9", "八", pdf_page=1),
            _meta(title, "zzz-early", "pdf:1:lines:1-7", "一", pdf_page=1),
            _meta(title, "mmm-p2", "pdf:2:lines:1-2", "二", pdf_page=2),
        ]
        con.executemany(
            "INSERT INTO chunks (id, metadata) VALUES (?, ?)",
            [(f"r{i}", json.dumps(m, ensure_ascii=False)) for i, m in enumerate(payload)],
        )
        con.commit()
        ids = [r["chunk_id"] for r in load_chunks(titles=[title], connection=con)]
        self.assertEqual(ids, ["zzz-early", "aaa-late", "mmm-p2"])

    def test_说明页包含范围与判定口径(self):
        wb = build_review_workbook(self._rows(), {"run_id": "run-x"})
        self.assertIn("填写说明", wb.sheetnames)
        text = "\n".join(
            str(c.value) for row in wb["填写说明"].iter_rows() for c in row if c.value
        )
        self.assertIn("判定口径", text)
        self.assertIn("run-x", text)
        self.assertIn("抽检", text)

    def test_导出的工作簿可被重新打开(self):
        buffer = io.BytesIO()
        build_review_workbook(self._rows(), {}).save(buffer)
        buffer.seek(0)
        reopened = load_workbook(buffer)
        self.assertEqual(reopened["待审核清单"].max_row, 4)


if __name__ == "__main__":
    unittest.main()
