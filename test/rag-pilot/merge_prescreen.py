"""合并 AI 预审结果，生成供人工确认的建议清单。

原则：
- 预审只是建议，**判定列留空**，必须由真实审核人选择并签名；
- 覆盖情况逐份文档统计，缺行直接报出来；
- 不修改原始审核清单，产出新文件。
"""

import argparse
import json
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

VALID_SUGGESTIONS = {"通过", "待修改", "驳回", "待补充原文"}

COLUMNS = [
    "序号",
    "片段ID（勿改）",
    "文档",
    "原文位置",
    "片段文本",
    "字数",
    "预审建议",
    "预审理由",
    "风险标记",
    "判定",
    "适用人群",
    "护理场景",
    "修改建议",
    "审核人",
    "审核日期",
]
WIDTHS = [6, 20, 26, 18, 70, 8, 12, 40, 28, 12, 14, 16, 24, 10, 12]


def normalize_entry(item, seq_map=None):
    """兼容两种写法：完整字段，或分批过程稿的简写（seq/s/p/c/f/r）。

    seq 是审核清单里的「序号」，可用它反查 chunk_id。
    """
    if not isinstance(item, dict):
        return None
    chunk_id = (item.get("chunk_id") or "").strip()
    if not chunk_id and seq_map is not None:
        try:
            chunk_id = seq_map.get(int(item.get("seq"))) or ""
        except (TypeError, ValueError):
            chunk_id = ""
    if not chunk_id:
        return None
    flags = item.get("risk_flags") or item.get("f") or []
    return {
        "chunk_id": chunk_id,
        "doc": item.get("doc", ""),
        "locator": item.get("locator", ""),
        "suggestion": item.get("suggestion") or item.get("s") or "",
        "population": item.get("population") or item.get("p") or "",
        "care_stage": item.get("care_stage") or item.get("c") or "",
        "risk_flags": list(flags) if isinstance(flags, (list, tuple)) else [],
        "reason": item.get("reason") or item.get("r") or "",
        "seq": item.get("seq"),
    }


def collect_entries(items, seq_map=None):
    """按 chunk_id 去重，后出现的覆盖先出现的。"""
    merged = {}
    for item in items:
        normalized = normalize_entry(item, seq_map)
        if not normalized:
            continue
        merged[normalized["chunk_id"]] = normalized
    return merged


def load_entries(review_dir, seq_map=None):
    items = []
    for path in sorted(Path(review_dir).rglob("*.jsonl")):
        for line in path.read_text(encoding="utf-8").splitlines():
            line = line.strip()
            if not line:
                continue
            try:
                items.append(json.loads(line))
            except ValueError:
                continue
    return collect_entries(items, seq_map)


def filter_to_sheet(entries, sheet_ids):
    """只保留当前清单里的片段；历史轮次遗留的 chunk_id 直接丢弃。"""
    return {cid: item for cid, item in entries.items() if cid in sheet_ids}


def coverage(entries, sheet_ids):
    """sheet_ids: {chunk_id: doc}。返回总数、已覆盖、缺失与逐文档统计。"""
    by_doc = {}
    missing = []
    for chunk_id, doc in sheet_ids.items():
        stat = by_doc.setdefault(doc, {"total": 0, "covered": 0, "missing": 0})
        stat["total"] += 1
        if chunk_id in entries:
            stat["covered"] += 1
        else:
            stat["missing"] += 1
            missing.append(chunk_id)
    return {
        "total": len(sheet_ids),
        "covered": len(sheet_ids) - len(missing),
        "missing": missing,
        "by_doc": by_doc,
    }


def build_suggestion_rows(items):
    rows = []
    for index, item in enumerate(items, start=1):
        rows.append(
            {
                "序号": index,
                "片段ID（勿改）": item.get("chunk_id", ""),
                "文档": item.get("doc", ""),
                "原文位置": item.get("locator", ""),
                "片段文本": item.get("text", ""),
                "字数": len(item.get("text", "")),
                "预审建议": item.get("suggestion", ""),
                "预审理由": item.get("reason", ""),
                "风险标记": "、".join(item.get("risk_flags") or []),
                "判定": "",
                "适用人群": item.get("population", ""),
                "护理场景": item.get("care_stage", ""),
                "修改建议": "",
                "审核人": "",
                "审核日期": "",
            }
        )
    return rows


def read_sheet(xlsx_path):
    workbook = load_workbook(xlsx_path)
    sheet = workbook["待审核清单"]
    headers = {cell.value: idx for idx, cell in enumerate(sheet[1], start=1)}

    def get(row, name):
        idx = headers.get(name)
        return sheet.cell(row=row, column=idx).value if idx else None

    records = []
    for row in range(2, sheet.max_row + 1):
        chunk_id = get(row, "片段ID（勿改）")
        if not chunk_id:
            continue
        records.append(
            {
                "seq": row - 1,
                "chunk_id": str(chunk_id).strip(),
                "doc": get(row, "文档") or "",
                "locator": get(row, "原文位置") or "",
                "text": get(row, "片段文本") or "",
                "section": get(row, "章节") or "",
            }
        )
    return records


def write_suggestion_workbook(rows, out_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "预审建议"

    header_fill = PatternFill("solid", fgColor="C08A7B")
    for col, (name, width) in enumerate(zip(COLUMNS, WIDTHS), start=1):
        cell = sheet.cell(row=1, column=col, value=name)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.column_dimensions[get_column_letter(col)].width = width
    sheet.freeze_panes = "C2"

    for row_index, row in enumerate(rows, start=2):
        for col, name in enumerate(COLUMNS, start=1):
            cell = sheet.cell(row=row_index, column=col, value=row.get(name, ""))
            cell.alignment = Alignment(vertical="top", wrap_text=col in (3, 4, 5, 8, 9, 13))
        sheet.row_dimensions[row_index].height = 54

    last = len(rows) + 1
    for column, options in (
        ("J", ("通过", "驳回", "待修改", "待补充原文")),
        ("K", ("早产儿", "早产儿+足月儿", "足月儿", "特定疾病人群", "不适用", "不确定")),
        ("L", ("住院期间", "出院当天", "出院后-居家", "随访", "门诊", "不适用", "不确定")),
    ):
        validation = DataValidation(
            type="list", formula1='"' + ",".join(options) + '"', allow_blank=True
        )
        sheet.add_data_validation(validation)
        validation.add(f"{column}2:{column}{last}")

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(out_path)
    return out_path


def main(argv=None):
    parser = argparse.ArgumentParser(description="合并 AI 预审结果")
    parser.add_argument("--review-dir", default="library/review")
    parser.add_argument("--sheet", required=True, help="原始审核清单 xlsx")
    parser.add_argument("--out", required=True, help="输出的预审建议 xlsx")
    parser.add_argument("--json-out", help="同时输出合并后的 json")
    args = parser.parse_args(argv)

    base = Path(__file__).resolve().parent
    review_dir = Path(args.review_dir)
    if not review_dir.is_absolute():
        review_dir = base / review_dir

    sheet_rows = read_sheet(args.sheet)
    seq_map = {row["seq"]: row["chunk_id"] for row in sheet_rows}
    entries = load_entries(review_dir, seq_map)
    sheet_ids = {row["chunk_id"]: row["doc"] for row in sheet_rows}
    before = len(entries)
    entries = filter_to_sheet(entries, sheet_ids)
    if before != len(entries):
        print(f"忽略 {before - len(entries)} 条历史轮次遗留的预审记录")
    report = coverage(entries, sheet_ids)

    print(f"预审记录 {len(entries)} 条，清单 {report['total']} 行，已覆盖 {report['covered']}，缺 {len(report['missing'])}")
    for doc, stat in report["by_doc"].items():
        print(f"  {stat['covered']:>4}/{stat['total']:<4} {doc}")

    merged_rows = []
    for index, row in enumerate(sheet_rows, start=1):
        item = entries.get(row["chunk_id"], {})
        suggestion = item.get("suggestion", "")
        merged_rows.append(
            {
                "序号": index,
                "片段ID（勿改）": row["chunk_id"],
                "文档": row["doc"],
                "原文位置": row["locator"],
                "片段文本": row["text"],
                "字数": len(row["text"] or ""),
                "预审建议": suggestion if suggestion in VALID_SUGGESTIONS else "",
                "预审理由": item.get("reason", ""),
                "风险标记": "、".join(item.get("risk_flags") or []),
                "判定": "",
                "适用人群": item.get("population", ""),
                "护理场景": item.get("care_stage", ""),
                "修改建议": "",
                "审核人": "",
                "审核日期": "",
            }
        )

    out_path = write_suggestion_workbook(merged_rows, args.out)
    print(f"已写出 {out_path}")

    if args.json_out:
        payload = {
            "covered": report["covered"],
            "total": report["total"],
            "missing": report["missing"],
            "entries": entries,
        }
        Path(args.json_out).write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"合并 json → {args.json_out}")

    return 0 if not report["missing"] else 5


if __name__ == "__main__":
    sys.exit(main())
