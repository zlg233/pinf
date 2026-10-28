"""导出"待审核清单"（Excel），供护理团队逐条审核。

输入：本地待审核索引快照（vectors.sqlite，chunk 元数据为 JSON）。
输出：.xlsx —— 主表"待审核清单" + 说明页"填写说明"。

设计约束：
- 只读索引，不修改任何入库产物；不改 review_status。
- 片段ID 原样导出，回填时用它定位，避免按行号对齐。
- 判定/人群/场景三列用下拉，减少自由文本造成的回填歧义。
"""

from __future__ import annotations

import argparse
import json
import sqlite3
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# ── 第一轮审核范围：6 份指南/共识（汇编《语料库（终版）》不在本轮） ──
REVIEW_SCOPE_TITLES = (
    "关于印发早产儿保健工作规范的通知",
    "早产儿保健服务指南",
    "NICU出院高危儿0～3岁生长发育随访管理技术的专家共识 (1)",
    "新生儿早期基本保健技术的临床实施建议",
    "特殊状态儿童预防接种(广东)专家共识",
    "新生儿常见呼吸道病毒感染预防的专家建议",
)

# 来自入库报告的逐文件质量事项，按标题前缀匹配
QUALITY_HINTS = {
    "关于印发早产儿保健工作规范的通知": "含表格（complex_table）；文本稀疏（sparse_text），可能漏段",
    "早产儿保健服务指南": "含表格（complex_table），表格行列可能错位",
    "NICU出院高危儿": "疑似双栏（possible_columns），读序需核对",
    "新生儿早期基本保健技术的临床实施建议": "疑似双栏（possible_columns），读序需核对",
    "特殊状态儿童预防接种": "疑似双栏（possible_columns），读序需核对",
    "新生儿常见呼吸道病毒感染预防的专家建议": "疑似双栏（possible_columns），读序需核对",
}

# 内容与其它文档逐字重复、本轮不单独审核的文档（实测：DOCX 附件 122 条全部包含在 PDF 通知 169 条里）
DUPLICATE_TITLES = {
    "早产儿保健服务指南": "关于印发早产儿保健工作规范的通知",
}

JUDGE_OPTIONS = ("通过", "驳回", "待修改", "待补充原文")
POPULATION_OPTIONS = ("早产儿", "早产儿+足月儿", "足月儿", "特定疾病人群", "不适用", "不确定")
STAGE_OPTIONS = ("住院期间", "出院当天", "出院后-居家", "随访", "门诊", "不适用", "不确定")

HEADERS = [
    "序号",
    "片段ID（勿改）",
    "文档",
    "类型",
    "原文位置",
    "页",
    "章节",
    "片段文本",
    "父段落",
    "字数",
    "质量提示",
    "判定",
    "适用人群",
    "护理场景",
    "修改建议",
    "审核人",
    "审核日期",
    "备注",
]

WIDTHS = [6, 20, 26, 8, 20, 6, 14, 70, 46, 8, 26, 12, 14, 16, 24, 10, 12, 18]


def _text_of(value) -> str:
    return "" if value is None else str(value)


def _doc_type(meta: dict) -> str:
    raw = _text_of(meta.get("raw_path")).lower()
    if raw.endswith(".pdf"):
        return "PDF"
    if raw.endswith(".docx"):
        return "DOCX"
    if raw.endswith(".doc"):
        return "DOC"
    return "未知"


def human_locator(meta: dict) -> str:
    """把机器 locator 转成护理团队看得懂的定位描述。"""
    locator = _text_of(meta.get("locator")).strip()
    if not locator:
        page = meta.get("pdf_page")
        return f"PDF 第 {page} 页" if page else "位置未知"

    parts = locator.split(":")
    if parts[0] == "pdf" and len(parts) >= 2:
        page = parts[1]
        lines = parts[3] if len(parts) >= 4 and parts[2] == "lines" else ""
        page_end = _text_of(meta.get("pdf_page_end"))
        page_text = f"PDF 第 {page} 页" if page_end in ("", page) else f"PDF 第 {page}-{page_end} 页"
        return f"{page_text} · 行 {lines}" if lines else page_text

    if parts[0] == "paragraph" and len(parts) >= 2:
        return f"DOCX 第 {parts[1]} 段"

    page = meta.get("pdf_page")
    return f"PDF 第 {page} 页" if page else locator


def _quality_hint(title: str) -> str:
    for prefix, hint in QUALITY_HINTS.items():
        if title.startswith(prefix):
            return hint
    return ""


def _line_start(locator: str) -> int:
    """从 locator 里取排序锚点：pdf:1:lines:8-9 → 8；paragraph:3 → 3。"""
    parts = locator.split(":")
    if len(parts) >= 4 and parts[2] == "lines":
        first = parts[3].split("-")[0]
        if first.isdigit():
            return int(first)
    if len(parts) >= 2 and parts[0] == "paragraph" and parts[1].isdigit():
        return int(parts[1])
    return 0


def _sort_key(index: int):
    def key(row: dict):
        page = row.get("pdf_page") or 0
        return (index, page, _line_start(row.get("locator", "")), row.get("_seq", 0))

    return key


def load_chunks(db_path=None, titles=(), connection=None, include_duplicates: bool = False) -> list[dict]:
    """按给定文档标题读片段；顺序为文档注册顺序 → 页码 → 定位。

    可传 db_path（只读打开）或已建立的 connection（便于测试用内存库）。
    默认跳过 DUPLICATE_TITLES 里与其它文档逐字重复的资料。
    """
    wanted = [
        title
        for title in titles
        if include_duplicates or title not in DUPLICATE_TITLES
    ]
    order = {title: idx for idx, title in enumerate(wanted)}
    own_con = connection is None
    if own_con:
        connection = sqlite3.connect(f"file:{Path(db_path).as_posix()}?mode=ro", uri=True)
    try:
        rows = connection.execute("SELECT rowid, metadata FROM chunks").fetchall()
    finally:
        if own_con:
            connection.close()

    picked: list[dict] = []
    for seq, (_, metadata) in enumerate(rows):
        meta = json.loads(metadata)
        title = _text_of(meta.get("title"))
        if title not in order:
            continue
        picked.append(
            {
                "title": title,
                "order": order[title],
                "_seq": seq,
                "chunk_id": _text_of(meta.get("chunk_id")),
                "locator": _text_of(meta.get("locator")),
                "section": _text_of(meta.get("section")),
                "pdf_page": meta.get("pdf_page"),
                "text": _text_of(meta.get("text")),
                "parent_text": _text_of(meta.get("parent_text")),
                "type": _doc_type(meta),
                "quality": _quality_hint(title),
                "document_sha256": _text_of(meta.get("document_sha256")),
            }
        )

    picked.sort(key=lambda row: _sort_key(row["order"])(row))
    return picked


def _style_header(ws, headers, widths) -> None:
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="C08A7B")
    thin = Side(style="thin", color="D9C8C0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col, (name, width) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
        ws.column_dimensions[get_column_letter(col)].width = width

    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
    ws.row_dimensions[1].height = 30


def _add_list_validation(ws, column_letter: str, last_row: int, options) -> None:
    dv = DataValidation(
        type="list",
        formula1='"' + ",".join(options) + '"',
        allow_blank=True,
        showDropDown=False,
    )
    dv.error = "请从下拉列表中选择"
    dv.errorTitle = "取值不在允许范围"
    dv.prompt = "选择一项"
    dv.promptTitle = "可选项"
    ws.add_data_validation(dv)
    dv.add(f"{column_letter}2:{column_letter}{last_row + 1}")


def build_review_workbook(rows: list[dict], meta: dict | None = None) -> Workbook:
    meta = meta or {}
    wb = Workbook()
    ws = wb.active
    ws.title = "待审核清单"
    _style_header(ws, HEADERS, WIDTHS)

    thin = Side(style="thin", color="E6DAD4")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for idx, row in enumerate(rows, start=1):
        excel_row = idx + 1
        values = [
            idx,
            row.get("chunk_id", ""),
            row.get("title", ""),
            row.get("type", ""),
            human_locator(row),
            row.get("pdf_page") or "",
            row.get("section", ""),
            row.get("text", ""),
            row.get("parent_text", ""),
            len(row.get("text", "")),
            row.get("quality", ""),
            "",
            "",
            "",
            "",
            "",
            "",
            "",
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=excel_row, column=col, value=value)
            cell.border = border
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=col in (3, 5, 7, 8, 9, 11, 15, 18),
                horizontal="center" if col in (1, 4, 6, 10, 12, 13, 14, 16, 17) else "left",
            )
        ws.row_dimensions[excel_row].height = 60

    last_row = len(rows) + 1
    _add_list_validation(ws, "L", last_row, JUDGE_OPTIONS)
    _add_list_validation(ws, "M", last_row, POPULATION_OPTIONS)
    _add_list_validation(ws, "N", last_row, STAGE_OPTIONS)

    _build_instructions_sheet(wb, rows, meta)
    return wb


def _build_instructions_sheet(wb: Workbook, rows: list[dict], meta: dict) -> None:
    ws = wb.create_sheet("填写说明")
    ws.column_dimensions["A"].width = 100

    counts: dict[str, int] = {}
    for row in rows:
        counts[row["title"]] = counts.get(row["title"], 0) + 1

    lines = [
        ("标题", "早产儿知识库 · 第一轮待审核清单（6 份指南/共识）"),
        ("", ""),
        ("导出信息", ""),
        ("运行快照", _text_of(meta.get("run_id")) or "-"),
        ("向量模型", _text_of(meta.get("model_id")) or "-"),
        ("维度", _text_of(meta.get("dimension")) or "-"),
        ("导出时间", _text_of(meta.get("exported_at")) or datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("导出条数", f"{len(rows)} 条"),
        ("", ""),
        ("本轮范围（按文档）", ""),
    ]
    for title, count in counts.items():
        lines.append((f"  {title}", f"{count} 条"))
    skipped = [t for t in DUPLICATE_TITLES if t not in counts]
    if skipped:
        lines.append(("", ""))
        lines.append(("重复内容（未单独导出）", ""))
        for title in skipped:
            lines.append((f"  {title}", f"与《{DUPLICATE_TITLES[title]}》逐字重复，随该文档一起审核"))
    lines += [
        ("", ""),
        ("判定口径", ""),
        ("  通过", "内容可直接用于回答用户，出处、数字、条件与例外都成立"),
        ("  驳回", "内容错误、过时、或与原文不符，且不值得修改"),
        ("  待修改", "方向可用，但需要改写/补条件，请在“修改建议”里写清怎么改"),
        ("  待补充原文", "原文解析缺失或错位，需要重新解析，交由技术处理"),
        ("  只要拿不准，不要选“通过”", "宁可不放行"),
        ("", ""),
        ("适用人群", " / ".join(POPULATION_OPTIONS)),
        ("护理场景", " / ".join(STAGE_OPTIONS)),
        ("", ""),
        ("抽检规则", "全部审完后，另找一人复核 10%；含数字/剂量/条件的片段 100% 复核"),
        ("回填方式", "把填写完成的文件交回，由技术同学运行导入脚本生成 approved 快照；不直接改线上数据库"),
        ("注意事项", "不要插入/删除行，不要修改“片段ID（勿改）”列；只填 L~R 列"),
        ("质量提示", "“质量提示”列来自自动解析报告，提示该文档可能存在表格错位或双栏读序问题，请重点核对"),
    ]

    for row_idx, (left, right) in enumerate(lines, start=1):
        ws.cell(row=row_idx, column=1, value=left).alignment = Alignment(
            vertical="top", wrap_text=True
        )
        ws.cell(row=row_idx, column=2, value=right).alignment = Alignment(
            vertical="top", wrap_text=True
        )
        if left in ("标题",):
            ws.cell(row=row_idx, column=1).font = Font(bold=True, size=14)
        elif left and not left.startswith("  "):
            ws.cell(row=row_idx, column=1).font = Font(bold=True)
    ws.column_dimensions["B"].width = 80


def resolve_index_path(explicit: str | None) -> Path:
    if explicit:
        return Path(explicit)
    base = Path(__file__).resolve().parent
    current = json.loads((base / "library/indexes/CURRENT.json").read_text(encoding="utf-8"))
    return base / "library" / current["index"]


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description="导出待审核清单 xlsx")
    parser.add_argument("--index", help="vectors.sqlite 路径，默认取 CURRENT.json")
    parser.add_argument("--out", help="输出 xlsx 路径")
    parser.add_argument("--titles", nargs="*", help="覆盖默认审核范围（文档标题）")
    parser.add_argument("--include-compiled", action="store_true", help="把汇编语料一起导出")
    parser.add_argument("--include-duplicates", action="store_true", help="把与其它文档重复的资料也导出")
    args = parser.parse_args(argv)

    index_path = resolve_index_path(args.index)
    if not index_path.exists():
        print(f"索引不存在：{index_path}", file=sys.stderr)
        return 2

    titles = args.titles
    if not titles:
        titles = list(REVIEW_SCOPE_TITLES)
        if args.include_compiled:
            titles.append("语料库（终版）")

    rows = load_chunks(index_path, titles, include_duplicates=args.include_duplicates)
    if not rows:
        print("没有匹配到任何片段，请检查文档标题是否与索引一致", file=sys.stderr)
        return 3

    current = json.loads((Path(__file__).resolve().parent / "library/indexes/CURRENT.json").read_text(encoding="utf-8"))
    meta = {
        "run_id": current.get("run_id"),
        "model_id": current.get("model_id"),
        "dimension": current.get("dimension"),
        "exported_at": datetime.now().strftime("%Y-%m-%d %H:%M"),
        "index_path": str(index_path),
    }

    wb = build_review_workbook(rows, meta)
    out_path = Path(args.out) if args.out else (
        Path(__file__).resolve().parent
        / "library/review"
        / f"round1-guides-{datetime.now().strftime('%Y%m%d-%H%M')}.xlsx"
    )
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)

    counts: dict[str, int] = {}
    for row in rows:
        counts[row["title"]] = counts.get(row["title"], 0) + 1
    print(f"已导出 {len(rows)} 条 → {out_path}")
    for title, count in counts.items():
        print(f"  {count:>4}  {title}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
