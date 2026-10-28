"""把审核清单（xlsx）回填成 approved 快照。

流程：读取已填写的审核表 → 校验 → 生成 decisions JSON（审计留痕）→
只把「通过」的片段写成新的 SQLite 快照（供后续导入 pgvector）。

硬约束：
- approved 为 0 时拒绝产出（退出码 4），不允许空快照上线；
- 校验失败按行报错（退出码 3），不部分写入；
- 只写通过项，pending/驳回/待修改绝不进入快照。
"""

from __future__ import annotations

import argparse
import json
import re
import sqlite3
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

DECISION_TO_STATUS = {
    "通过": "approved",
    "驳回": "rejected",
    "待修改": "needs_revision",
    "待补充原文": "needs_source",
}
STATUS_ORDER = ("approved", "rejected", "needs_revision", "needs_source", "unjudged")

COL_CHUNK = "片段ID（勿改）"
COL_DOC = "文档"
COL_DECISION = "判定"
COL_POPULATION = "适用人群"
COL_STAGE = "护理场景"
COL_NOTE = "修改建议"
COL_REVIEWER = "审核人"
COL_DATE = "审核日期"
COL_TEXT = "片段文本"

REQUIRED_HEADERS = (COL_CHUNK, COL_DECISION, COL_POPULATION, COL_STAGE, COL_REVIEWER, COL_DATE)

# 机器/AI 预审不能充当人工签字
AI_REVIEWER_PATTERN = re.compile(r"^\s*(ai|a\.i\.|机器|预审|gpt|claude|machine|robot)", re.IGNORECASE)


def _clean(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    return str(value).strip()


def _normalize_date(value) -> tuple[str, bool]:
    """返回 (规范日期, 是否合法)。空值视为合法（回填时补当天）。"""
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d"), True
    text = _clean(value)
    if not text:
        return "", True
    for fmt in ("%Y-%m-%d", "%Y/%m/%d"):
        try:
            return datetime.strptime(text, fmt).strftime("%Y-%m-%d"), True
        except ValueError:
            continue
    return text, False


def _header_map(ws) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for idx, cell in enumerate(ws[1], start=1):
        name = _clean(cell.value)
        if name:
            mapping[name] = idx
    missing = [h for h in REQUIRED_HEADERS if h not in mapping]
    if missing:
        raise ValueError("审核表缺少必需列：" + "、".join(missing))
    return mapping


def parse_review_workbook(workbook) -> list[dict]:
    """按表头名字取列，避免列顺序变化导致错位。"""
    ws = workbook["待审核清单"]
    headers = _header_map(ws)

    def get(row_cells, name):
        idx = headers.get(name)
        return row_cells[idx - 1].value if idx else None

    rows: list[dict] = []
    for row_cells in ws.iter_rows(min_row=2):
        chunk_id = _clean(get(row_cells, COL_CHUNK))
        if not chunk_id:
            continue
        reviewed_date, date_ok = _normalize_date(get(row_cells, COL_DATE))
        rows.append(
            {
                "row": row_cells[0].row,
                "chunk_id": chunk_id,
                "title": _clean(get(row_cells, COL_DOC)),
                "text": _clean(get(row_cells, COL_TEXT))[:200],
                "decision": _clean(get(row_cells, COL_DECISION)),
                "population": _clean(get(row_cells, COL_POPULATION)),
                "care_stage": _clean(get(row_cells, COL_STAGE)),
                "reviewer": _clean(get(row_cells, COL_REVIEWER)),
                "reviewed_date": reviewed_date,
                "_date_ok": date_ok,
                "note": _clean(get(row_cells, COL_NOTE)),
            }
        )
    return rows


def validate_rows(rows: list[dict], known_ids: set[str]) -> list[dict]:
    errors: list[dict] = []
    seen: dict[str, int] = {}

    for item in rows:
        decision = item["decision"]
        if not decision:
            continue
        row_no = item["row"]

        if decision not in DECISION_TO_STATUS:
            errors.append(
                {
                    "row": row_no,
                    "code": "invalid_decision",
                    "message": f"判定“{decision}”不在允许值内（通过/驳回/待修改/待补充原文）",
                }
            )
            continue

        if item["chunk_id"] not in known_ids:
            errors.append(
                {
                    "row": row_no,
                    "code": "unknown_chunk",
                    "message": f"片段ID {item['chunk_id']} 不在当前索引中",
                }
            )

        if not item["reviewer"]:
            errors.append({"row": row_no, "code": "missing_reviewer", "message": "已做判定但未填审核人"})
        elif AI_REVIEWER_PATTERN.match(item["reviewer"]):
            # 机器预审不能替代人工签字，否则 approved 的来源不可追责
            errors.append(
                {
                    "row": row_no,
                    "code": "ai_reviewer",
                    "message": f"审核人“{item['reviewer']}”看起来是机器预审，必须填真实审核人",
                }
            )

        if decision == "通过":
            for field, label in (("population", "适用人群"), ("care_stage", "护理场景")):
                if not item[field]:
                    errors.append(
                        {"row": row_no, "code": "missing_meta", "message": f"判定为通过但未填{label}"}
                    )

        _, date_ok = _normalize_date(item.get("reviewed_date"))
        if not date_ok:
            errors.append(
                {
                    "row": row_no,
                    "code": "bad_date",
                    "message": f"审核日期“{item['reviewed_date']}”不是 YYYY-MM-DD",
                }
            )
        if item["chunk_id"] in seen:
            errors.append(
                {
                    "row": row_no,
                    "code": "duplicate_row",
                    "message": f"片段ID重复，首次出现在第 {seen[item['chunk_id']]} 行",
                }
            )
        else:
            seen[item["chunk_id"]] = row_no

    return errors


def build_decisions(rows: list[dict], meta: dict | None = None) -> dict:
    meta = meta or {}
    counts = {status: 0 for status in STATUS_ORDER}
    items: list[dict] = []

    for item in rows:
        status = DECISION_TO_STATUS.get(item["decision"], "unjudged")
        counts[status] += 1
        items.append(
            {
                "chunk_id": item["chunk_id"],
                "title": item.get("title", ""),
                "decision": item["decision"],
                "status": status,
                "population": item["population"],
                "care_stage": item["care_stage"],
                "reviewer": item["reviewer"],
                "reviewed_date": item["reviewed_date"],
                "note": item["note"],
            }
        )

    return {
        "created_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "run_id": meta.get("run_id"),
        "source_file": meta.get("source_file"),
        "counts": counts,
        "items": items,
    }


def decide_exit_code(counts: dict, error_count: int) -> int:
    if error_count:
        return 3
    if counts.get("approved", 0) <= 0:
        return 4
    return 0


def write_approved_snapshot(source: sqlite3.Connection, target: sqlite3.Connection, decisions: dict) -> int:
    """把 approved 项复制到新快照，并写入审核字段。返回写入条数。"""
    target.execute("CREATE TABLE IF NOT EXISTS chunks (id TEXT, metadata TEXT, vector BLOB)")
    written = 0
    for item in decisions.get("items", []):
        if item.get("status") != "approved":
            continue
        row = source.execute(
            "SELECT metadata, vector FROM chunks WHERE json_extract(metadata, '$.chunk_id') = ?",
            (item["chunk_id"],),
        ).fetchone()
        if row is None:
            continue
        metadata = json.loads(row[0])
        metadata.update(
            {
                "review_status": "approved",
                "reviewed_by": item["reviewer"],
                "reviewed_at": item["reviewed_date"] or datetime.now().strftime("%Y-%m-%d"),
                "review_note": item["note"],
                "population": item["population"],
                "care_stage": item["care_stage"],
            }
        )
        target.execute(
            "INSERT INTO chunks (id, metadata, vector) VALUES (?, ?, ?)",
            (item["chunk_id"], json.dumps(metadata, ensure_ascii=False), row[1]),
        )
        written += 1
    target.commit()
    return written


def _load_index_ids(index_path: Path) -> set[str]:
    con = sqlite3.connect(f"file:{index_path.as_posix()}?mode=ro", uri=True)
    try:
        return {
            json.loads(metadata).get("chunk_id")
            for (metadata,) in con.execute("SELECT metadata FROM chunks")
        }
    finally:
        con.close()


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description="审核回填：xlsx → approved 快照")
    parser.add_argument("--input", required=True, help="已填写的审核清单 xlsx")
    parser.add_argument("--index", help="源索引 vectors.sqlite，默认取 CURRENT.json")
    parser.add_argument("--out-dir", help="输出目录，默认 library/review")
    parser.add_argument("--dry-run", action="store_true", help="只校验与统计，不写任何文件")
    args = parser.parse_args(argv)

    base = Path(__file__).resolve().parent
    current = json.loads((base / "library/indexes/CURRENT.json").read_text(encoding="utf-8"))
    index_path = Path(args.index) if args.index else base / "library" / current["index"]
    out_dir = Path(args.out_dir) if args.out_dir else base / "library/review"

    if not index_path.exists():
        print(f"索引不存在：{index_path}", file=sys.stderr)
        return 2
    input_path = Path(args.input)
    if not input_path.exists():
        print(f"审核表不存在：{input_path}", file=sys.stderr)
        return 2

    workbook = load_workbook(input_path, data_only=True)
    rows = parse_review_workbook(workbook)
    known_ids = _load_index_ids(index_path)
    errors = validate_rows(rows, known_ids)

    decisions = build_decisions(
        rows, {"run_id": current.get("run_id"), "source_file": input_path.name}
    )
    counts = decisions["counts"]

    print(f"读取 {len(rows)} 行：通过 {counts['approved']} / 驳回 {counts['rejected']} / "
          f"待修改 {counts['needs_revision']} / 待补充原文 {counts['needs_source']} / 未判定 {counts['unjudged']}")

    if errors:
        print(f"校验失败 {len(errors)} 处：", file=sys.stderr)
        for err in errors[:20]:
            print(f"  第 {err['row']} 行 [{err['code']}] {err['message']}", file=sys.stderr)
        return 3

    if args.dry_run:
        print("dry-run：未写任何文件")
        return 0

    exit_code = decide_exit_code(counts, len(errors))
    if exit_code == 4:
        print("没有任何片段通过审核，拒绝生成快照（避免空知识库上线）", file=sys.stderr)
        return 4

    stamp = datetime.now().strftime("%Y%m%d-%H%M")
    out_dir.mkdir(parents=True, exist_ok=True)
    decisions_path = out_dir / f"decisions-{stamp}.json"
    decisions_path.write_text(
        json.dumps(decisions, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    snapshot_dir = base / "library/indexes" / f"approved-{stamp}"
    snapshot_dir.mkdir(parents=True, exist_ok=True)
    snapshot_path = snapshot_dir / "vectors.sqlite"
    source = sqlite3.connect(f"file:{index_path.as_posix()}?mode=ro", uri=True)
    target = sqlite3.connect(snapshot_path)
    try:
        written = write_approved_snapshot(source, target, decisions)
    finally:
        source.close()
        target.close()

    print(f"已写入 approved 快照 {written} 条 → {snapshot_path}")
    print(f"审计记录 → {decisions_path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
